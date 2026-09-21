"""Core logic: storage, matching, tracking. No MCP code here, so it is easy to test."""
import json
import os
import sqlite3
from collections import Counter
from contextlib import contextmanager
from datetime import date, timedelta
from pathlib import Path

from .skills import extract_skills, work_pass_signals

VALID_STATUSES = ["saved", "applied", "interviewing", "offer", "rejected", "withdrawn"]
ACTIVE_STATUSES = ("applied", "interviewing")

SCHEMA = """
CREATE TABLE IF NOT EXISTS resume (
    id INTEGER PRIMARY KEY CHECK (id = 1),
    text TEXT NOT NULL,
    skills TEXT NOT NULL,
    updated TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS jobs (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    title TEXT NOT NULL,
    company TEXT NOT NULL,
    description TEXT NOT NULL,
    skills TEXT NOT NULL,
    source TEXT NOT NULL,
    url TEXT NOT NULL,
    created TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS applications (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    job_id INTEGER NOT NULL UNIQUE REFERENCES jobs(id),
    status TEXT NOT NULL,
    applied_date TEXT,
    followup_date TEXT,
    notes TEXT NOT NULL DEFAULT ''
);
"""


def _db_path() -> Path:
    default = Path(__file__).resolve().parent.parent / "data" / "jobhunt.db"
    return Path(os.environ.get("JOBHUNT_DB", default))


@contextmanager
def _conn():
    path = _db_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row
    conn.executescript(SCHEMA)
    try:
        yield conn
        conn.commit()
    finally:
        conn.close()


def _today() -> date:
    return date.today()


# ---------- resume ----------

def save_resume(text: str) -> dict:
    text = text.strip()
    if len(text) < 30:
        raise ValueError("Resume text looks too short. Paste the full resume text.")
    skills = extract_skills(text)
    with _conn() as c:
        c.execute(
            "INSERT INTO resume (id, text, skills, updated) VALUES (1, ?, ?, ?) "
            "ON CONFLICT(id) DO UPDATE SET text=excluded.text, skills=excluded.skills, updated=excluded.updated",
            (text, json.dumps(skills), _today().isoformat()),
        )
    return {"saved": True, "skills_found": skills, "skill_count": len(skills)}


def get_resume() -> dict:
    with _conn() as c:
        row = c.execute("SELECT * FROM resume WHERE id = 1").fetchone()
    if not row:
        raise ValueError("No resume saved yet. Call save_resume first.")
    return {"text": row["text"], "skills": json.loads(row["skills"]), "updated": row["updated"]}
def _data_dir() -> Path:
    default = Path(__file__).resolve().parent.parent / "data"
    return Path(os.environ.get("JOBHUNT_DATA_DIR", default))


def _read_file_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    raise ValueError("Unsupported file type. Use .pdf, .docx, .txt or .md.")


def load_resume_file(filename: str) -> dict:
    """Read a resume from the data folder (pdf, docx, txt, md) and save it."""
    data_dir = _data_dir().resolve()
    path = (data_dir / filename).resolve()
    if data_dir not in path.parents:
        raise ValueError("The file must be inside the data folder.")
    if not path.is_file():
        raise ValueError(f"File not found in the data folder: {filename}")
    text = _read_file_text(path).strip()
    if len(text) < 200:
        raise ValueError(
            "Very little text could be read. The file may be a scanned image or have an unusual layout. "
            "Try a simple single-column version, or paste the text instead."
        )
    result = save_resume(text)
    result.update({"file": path.name, "characters_read": len(text)})
    return result

# ---------- jobs ----------

def add_job(title: str, company: str, description: str, url: str = "", source: str = "manual") -> dict:
    description = description.strip()
    if len(description) < 30:
        raise ValueError("Job description looks too short. Paste the full description.")
    skills = extract_skills(description + " " + title)
    with _conn() as c:
        cur = c.execute(
            "INSERT INTO jobs (title, company, description, skills, source, url, created) VALUES (?,?,?,?,?,?,?)",
            (title.strip(), company.strip(), description, json.dumps(skills), source, url, _today().isoformat()),
        )
        job_id = cur.lastrowid
        c.execute("INSERT INTO applications (job_id, status) VALUES (?, 'saved')", (job_id,))
    return {
        "job_id": job_id,
        "title": title,
        "company": company,
        "skills_required": skills,
        "work_pass_signals": work_pass_signals(description),
    }


def list_jobs() -> list[dict]:
    with _conn() as c:
        rows = c.execute(
            "SELECT j.id, j.title, j.company, j.created, a.status, a.followup_date "
            "FROM jobs j JOIN applications a ON a.job_id = j.id ORDER BY j.id"
        ).fetchall()
    return [dict(r) for r in rows]


def get_job(job_id: int) -> dict:
    with _conn() as c:
        row = c.execute("SELECT * FROM jobs WHERE id = ?", (job_id,)).fetchone()
    if not row:
        raise ValueError(f"No job with id {job_id}.")
    job = dict(row)
    job["skills"] = json.loads(job["skills"])
    return job

def delete_job(job_id: int, confirm: bool = False) -> dict:
    """Delete a saved job and its tracking record. Needs confirm=True as a safety check."""
    if not confirm:
        raise ValueError("Deletion not done. Call again with confirm=True after the user clearly agrees.")
    job = get_job(job_id)  # raises if the job does not exist
    with _conn() as c:
        c.execute("DELETE FROM applications WHERE job_id = ?", (job_id,))
        c.execute("DELETE FROM jobs WHERE id = ?", (job_id,))
    return {"deleted": True, "job_id": job_id, "title": job["title"], "company": job["company"]}

# ---------- matching ----------

def match_resume_to_job(job_id: int) -> dict:
    resume = get_resume()
    job = get_job(job_id)
    have = set(resume["skills"])
    need = set(job["skills"])
    matched = sorted(need & have)
    missing = sorted(need - have)
    score = round(100 * len(matched) / len(need)) if need else None
    return {
        "job_id": job_id,
        "title": job["title"],
        "company": job["company"],
        "match_score_percent": score,
        "matched_skills": matched,
        "missing_skills": missing,
        "work_pass_signals": work_pass_signals(job["description"]),
        "note": (
            "Score counts only skills from the built-in skill list, so treat it as a rough guide. "
            "Work-pass signals are hints from the posting text; confirm eligibility with the employer."
        ),
    }


def get_tailoring_context(job_id: int) -> dict:
    """Everything an LLM needs to suggest resume edits, plus the honesty rule."""
    match = match_resume_to_job(job_id)
    return {
        "resume_text": get_resume()["text"],
        "job_title": match["title"],
        "company": match["company"],
        "job_description": get_job(job_id)["description"],
        "matched_skills": match["matched_skills"],
        "missing_skills": match["missing_skills"],
        "rules": [
            "Only reword or reorder experience that already appears in resume_text.",
            "Never claim a missing skill as experience. Suggest learning it or a project instead.",
            "Keep each bullet factual, specific, and under two lines.",
        ],
    }


# ---------- tracking ----------

def track_application(job_id: int, status: str, notes: str = "", followup_in_days: int = 7) -> dict:
    status = status.lower().strip()
    if status not in VALID_STATUSES:
        raise ValueError(f"Status must be one of {VALID_STATUSES}.")
    get_job(job_id)  # raises if missing
    today = _today()
    followup = (today + timedelta(days=followup_in_days)).isoformat() if status in ACTIVE_STATUSES else None
    with _conn() as c:
        row = c.execute("SELECT applied_date, notes FROM applications WHERE job_id = ?", (job_id,)).fetchone()
        applied = row["applied_date"] if row else None
        if status == "applied" and not applied:
            applied = today.isoformat()
        old_notes = row["notes"] if row else ""
        merged_notes = (old_notes + "\n" + notes).strip() if notes else old_notes
        c.execute(
            "INSERT INTO applications (job_id, status, applied_date, followup_date, notes) VALUES (?,?,?,?,?) "
            "ON CONFLICT(job_id) DO UPDATE SET status=excluded.status, applied_date=excluded.applied_date, "
            "followup_date=excluded.followup_date, notes=excluded.notes",
            (job_id, status, applied, followup, merged_notes),
        )
    return {"job_id": job_id, "status": status, "applied_date": applied, "followup_date": followup}


def get_followups_due() -> list[dict]:
    with _conn() as c:
        rows = c.execute(
            "SELECT j.id AS job_id, j.title, j.company, a.status, a.followup_date "
            "FROM applications a JOIN jobs j ON j.id = a.job_id "
            "WHERE a.followup_date IS NOT NULL AND a.followup_date <= ? AND a.status IN ('applied','interviewing') "
            "ORDER BY a.followup_date",
            (_today().isoformat(),),
        ).fetchall()
    return [dict(r) for r in rows]


def applications_summary() -> dict:
    jobs = list_jobs()
    counts = Counter(j["status"] for j in jobs)
    return {"total": len(jobs), "by_status": dict(counts), "jobs": jobs}


# ---------- market analysis ----------

def analyze_market_skills(top_n: int = 15) -> dict:
    with _conn() as c:
        rows = c.execute("SELECT skills FROM jobs").fetchall()
    if not rows:
        raise ValueError("No jobs saved yet. Add some jobs first.")
    counts: Counter = Counter()
    for r in rows:
        counts.update(json.loads(r["skills"]))
    try:
        have = set(get_resume()["skills"])
    except ValueError:
        have = set()
    total = len(rows)
    ranked = [
        {"skill": s, "jobs": n, "percent_of_jobs": round(100 * n / total), "you_have_it": s in have}
        for s, n in counts.most_common(top_n)
    ]
    return {
        "jobs_analyzed": total,
        "top_skills": ranked,
        "skills_to_learn_next": [r["skill"] for r in ranked if not r["you_have_it"]][:5],
    }
